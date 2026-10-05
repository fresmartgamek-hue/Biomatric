import base64
import io
import json
import math
import os
import sys
from datetime import datetime, time, timedelta
from flask import (
    Flask,
    render_template_string,
    request,
    Response,
    send_file,
    session,
    redirect,
    url_for,
    flash,
    send_from_directory,
)
from github import Github
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from werkzeug.utils import secure_filename
from zk import ZK, const

app = Flask(__name__)

# Secret configurations using Environment Variables with fallbacks
app.secret_key = os.getenv('SECRET_KEY', 'gamek_fresmart_secret_key_sonu')

UPLOAD_FOLDER = os.getenv('UPLOAD_FOLDER', 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

ALLOWED_EXTENSIONS = {'pdf', 'png', 'jpg', 'jpeg', 'doc', 'docx'}

MACHINE_IP = os.getenv('MACHINE_IP', '192.168.1.153')
PORT = int(os.getenv('MACHINE_PORT', 4370))

# GitHub Configurations via Environment Variables
GITHUB_TOKEN = os.getenv('GITHUB_TOKEN')
GITHUB_REPO_NAME = os.getenv('GITHUB_REPO_NAME', 'fresmartgamek-hue/Biometric')
GITHUB_BRANCH = os.getenv('GITHUB_BRANCH', 'main')

# Persistent Storage Files
LEAVE_JSON_FILE = 'leave_requests.json'
LEAVE_EXCEL_FILE = 'leave_records.xlsx'
CUSTOM_PUNCHES_FILE = 'custom_punches.json'
CUSTOM_OFFS_FILE = 'custom_weekly_offs.json'

# Global storage for synced biometric logs
SYNCED_ATTENDANCE_LOGS = []
LAST_DEVICE_SYNC_TIME = None

MASTER_EMPLOYEES = {
    'NWC2981': {
        'name': 'ANTONIO JOSE BANDOLA',
        'off': 'SUNDAY',
        'dept': 'ADMIN - MANAGER',
        'shift': 'morning',
    },
    'NWC3127': {
        'name': 'MATEUS ANTONIO DA COSTA BALMIRO',
        'off': 'FRIDAY',
        'dept': 'ADMIN - MANAGER',
        'shift': 'morning',
    },
    'NWC1525': {
        'name': 'ETY JOSÉ BANDUA MONTEIRO',
        'off': 'MONDAY',
        'dept': 'ADMIN - MANAGER',
        'shift': 'morning',
    },
    'NWC8328': {
        'name': 'JOAO MATIAS DOMINGOS',
        'off': 'SATURDAY',
        'dept': 'ADMIN - CCTV',
        'shift': 'morning',
    },
    'NWC6661': {
        'name': 'FRANCISCO MUNDELE CHIVELA',
        'off': 'WEDNESDAY',
        'dept': 'ADMIN - CCTV',
        'shift': 'morning',
    },
    'NWC1553': {
        'name': 'TIAGO SANDALA CHISSANHA',
        'off': 'SUNDAY',
        'dept': 'ADMIN - AUDITOR',
        'shift': 'morning',
    },
    'NWC8350': {
        'name': 'LOLIVALDO ALBERTO MADEIRA',
        'off': 'SUNDAY',
        'dept': 'ADMIN - EDP',
        'shift': 'morning',
    },
    'NWC5187': {
        'name': 'VICTOR NSOSI JOAO',
        'off': 'MONDAY',
        'dept': 'CASH - HEAD',
        'shift': 'morning',
    },
    'NWC1168': {
        'name': 'ADELIA MBALOMBO CHIPEPI',
        'off': 'SUNDAY',
        'dept': 'CASH - HEAD',
        'shift': 'morning',
    },
    'NWC2652': {
        'name': 'DULCE DOROTEIA GARCIA LUSITANO',
        'off': 'MONDAY',
        'dept': 'CASH - HEAD',
        'shift': 'morning',
    },
    'NWC3381': {
        'name': 'ANDRE DE JESUS NGOLA JOSE',
        'off': 'TUESDAY',
        'dept': 'CASH',
        'shift': 'morning',
    },
    'NWC1983': {
        'name': 'PATRICIA SOLANGE FRANCISCO',
        'off': 'THURSDAY',
        'dept': 'CASH',
        'shift': 'second',
    },
    'NWC8364': {
        'name': 'DIELUMBAKA AUGUSTO',
        'off': 'WEDNESDAY',
        'dept': 'CASH',
        'shift': 'second',
    },
    'NWC2788': {
        'name': 'INES NACHINGOLO FELICIANO NAMBELO',
        'off': 'TUESDAY',
        'dept': 'CASH',
        'shift': 'morning',
    },
    'NWC1010': {
        'name': 'TERESA PEDRO LEAO',
        'off': 'FRIDAY',
        'dept': 'CASH',
        'shift': 'morning',
    },
    'NWC6638': {
        'name': 'CLAUDIO JANUARIO MANUEL AVELINO',
        'off': 'SUNDAY',
        'dept': 'TALHO',
        'shift': 'morning',
    },
    'NWC5830': {
        'name': 'REGINA DE FATIMA VIDAL',
        'off': 'MONDAY',
        'dept': 'TALHO',
        'shift': 'morning',
    },
    'NWC5529': {
        'name': 'ALEXANDRE LUIS CORREIA',
        'off': 'FRIDAY',
        'dept': 'TALHO',
        'shift': 'morning',
    },
    'NWC5713': {
        'name': 'ROSA GARNEIRA BUMBA',
        'off': 'WEDNESDAY',
        'dept': 'TALHO',
        'shift': 'morning',
    },
    'NWC5396': {
        'name': 'COSTA BEBIANO HEBO',
        'off': 'THURSDAY',
        'dept': 'TALHO',
        'shift': 'morning',
    },
    'NWC8361': {
        'name': 'AGOSTINHO JOAQUIM KUANGO DA COSTA',
        'off': 'SUNDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC2300': {
        'name': 'JOANA CARDOSO JOAQUIM AFONSO',
        'off': 'MONDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC5168': {
        'name': 'JOAO NVUNDA DALA',
        'off': 'THURSDAY',
        'dept': 'F & V',
        'shift': 'morning',
    },
    'NWC3711': {
        'name': 'ANGELA MARIA BUMBA',
        'off': 'FRIDAY',
        'dept': 'F & V',
        'shift': 'morning',
    },
    'NWC5186': {
        'name': 'HELIA DOMINGOS DE CARVALHO',
        'off': 'WEDNESDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC6702': {
        'name': 'DOMINGOS GAMA PEREIRA',
        'off': 'FRIDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC3596': {
        'name': 'ALDAIR FERNANDES FERREIRA',
        'off': 'TUESDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC2757': {
        'name': 'JOSEFA KUELUNGA MUASSOKA',
        'off': 'THURSDAY',
        'dept': 'SECU',
        'shift': 'morning',
    },
    'NWC4554': {
        'name': 'DOMINGOS ANTONIO FERNANDO',
        'off': 'THURSDAY',
        'dept': 'CASH',
        'shift': 'morning',
    },
    'NWC2624': {
        'name': 'CECILIA JORGE FAMOSO',
        'off': 'FRIDAY',
        'dept': 'CASH',
        'shift': 'morning',
    },
    'NWC3318': {
        'name': 'JOSE MANUEL KAZOLA',
        'off': 'SUNDAY',
        'dept': 'FRESCO',
        'shift': 'morning',
    },
    'NWC7347': {
        'name': 'ARMANDO CHICOVO SAMBA',
        'off': 'TUESDAY',
        'dept': 'FRESCO',
        'shift': 'morning',
    },
    'NWC8362': {
        'name': 'ARAUJO PAULOMENDES',
        'off': 'FRIDAY',
        'dept': 'STOCK',
        'shift': 'morning',
    },
    'NWC6715': {
        'name': 'RIBEIRO ANTONIO FRANCISCO',
        'off': 'THURSDAY',
        'dept': 'STOCK',
        'shift': 'morning',
    },
    'NWC6444': {
        'name': 'HENRIQUES BRANDAO',
        'off': 'WEDNESDAY',
        'dept': 'STOCK',
        'shift': 'morning',
    },
}

EMPLOYEE_OVERRIDES = {
    '8364': {'code': 'NWC8364', 'name': 'DIELUMBAKA AUGUSTO'},
    '1': {'code': 'NWC8350', 'name': 'LOLIVALDO ALBERTO MADEIRA'},
    '8362': {'code': 'NWC8362', 'name': 'ARAUJO PAULOMENDES'},
    '6661': {'code': 'NWC6661', 'name': 'FRANCISCO MUNDELE CHIVELA'},
}

def load_json_file(filepath):
  if os.path.exists(filepath):
    try:
      with open(filepath, 'r', encoding='utf-8') as f:
        return json.load(f)
    except:
      return {}
  return {}

def save_json_file(filepath, data):
  try:
    with open(filepath, 'w', encoding='utf-8') as f:
      json.dump(data, f, indent=4, ensure_ascii=False)
  except Exception as e:
    print(f"Error saving {filepath}: {e}")

CUSTOM_PUNCHES = load_json_file(CUSTOM_PUNCHES_FILE)
CUSTOM_OFFS = load_json_file(CUSTOM_OFFS_FILE)

TRANSLATIONS = {
    'en': {
        'dashboard': 'Dashboard',
        'projects': 'Projects & Tasks',
        'calendar': 'Calendar & Rota',
        'leave_mgmt': 'Leave Management',
        'settings': 'Settings',
        'performance': 'Performance',
        'payroll': 'Payroll & Reports',
        'roster': 'Employees Roster',
        'announcements': 'Announcements',
        'biometric_active': 'Biometric live tracking active for Gamek Fresmart Express LM11.',
        'good_day': 'Good day',
        'online': 'Online',
        'offline': 'Offline',
        'sync': 'Sync',
        'logout': 'Logout',
        'shutdown': 'Shutdown',
        'present': 'Present',
        'absent': 'Absent',
        'medical_leave': 'Medical/Leave',
        'weekly_off': 'Weekly Off',
        'late_arrival': 'Late Arrival',
        'mispunches': 'Mis-Punches',
        'total_hrs': 'Total Hours',
        'start_date': 'Start Date',
        'end_date': 'End Date',
        'emp_filter': 'Employee Filter',
        'all_personnel': '-- All Personnel --',
        'export': 'Export',
        'shift_hint': 'Standard Duty: 7h w/o lunch (8h w/ lunch). Overtime counts if extra >= 45 mins.',
        'search_placeholder': '🔍 Search employee name, ID or department...',
        'sr': 'Sr. No.',
        'date': 'Date',
        'id': 'ID',
        'emp_name': 'Employee Name',
        'dept': 'Dept',
        'store_in': 'Store In',
        'lunch_out': 'Lunch Out',
        'lunch_in': 'Lunch In',
        'out_time': 'Out Time',
        'total_lunch': 'Total Lunch',
        'working_hours': 'Working Hours',
        'overtime': 'Total Hora Extra',
        'status': 'Status',
        'total_summary': 'Total Summary:',
        'close': 'Close',
        'cancel': 'Cancel',
        'apply_leave': 'Apply for Leave Request',
        'leave_type': 'Leave Type (Required)',
        'support_doc': 'Supporting Document Upload (Optional - PDF/Image)',
        'submit_req': 'Submit Request',
        'leave_history': 'Leave History Archive',
        'action': 'Action',
        'approve': 'Accept',
        'reject': 'Reject',
    },
    'pt': {
        'dashboard': 'Painel',
        'projects': 'Projetos e Tarefas',
        'calendar': 'Calendário e Escala',
        'leave_mgmt': 'Gestão de Licenças',
        'settings': 'Definições',
        'performance': 'Desempenho',
        'payroll': 'Folha e Relatórios',
        'roster': 'Lista de Funcionários',
        'announcements': 'Avisos',
        'biometric_active': 'Rastreio biométrico ao vivo ativo para Gamek Fresmart Express LM11.',
        'good_day': 'Bom dia',
        'online': 'Online',
        'offline': 'Offline',
        'sync': 'Sinc',
        'logout': 'Sair',
        'shutdown': 'Desligar',
        'present': 'Presente',
        'absent': 'Faltou',
        'medical_leave': 'Baixa/Licença',
        'weekly_off': 'Folga Semanal',
        'late_arrival': 'Atrasado',
        'mispunches': 'Marcação Falhada',
        'total_hrs': 'Total de Horas',
        'start_date': 'Data Inicial',
        'end_date': 'Data Final',
        'emp_filter': 'Filtro de Funcionários',
        'all_personnel': '-- Todo o Pessoal --',
        'export': 'Exportar',
        'shift_hint': 'Turno Padrão: 7h s/ almoço (8h c/ almoço). Hora extra se extra >= 45 mins.',
        'search_placeholder': '🔍 Pesquisar nome, ID ou departamento...',
        'sr': 'Nº',
        'date': 'Data',
        'id': 'ID',
        'emp_name': 'Nome do Funcionário',
        'dept': 'Dept',
        'store_in': 'Entrada',
        'lunch_out': 'Saída Almoço',
        'lunch_in': 'Entrada Almoço',
        'out_time': 'Saída',
        'total_lunch': 'Total Almoço',
        'working_hours': 'Horas de Trabalho',
        'overtime': 'Total Hora Extra',
        'status': 'Estado',
        'total_summary': 'Resumo Total:',
        'close': 'Fechar',
        'cancel': 'Cancelar',
        'apply_leave': 'Solicitar Licença',
        'leave_type': 'Tipo de Licença (Obrigatório)',
        'support_doc': 'Documento de Suporte (Opcional - PDF/Imagem)',
        'submit_req': 'Submeter Pedido',
        'leave_history': 'Arquivo de Histórico de Licenças',
        'action': 'Ação',
        'approve': 'Aceitar',
        'reject': 'Rejeitar',
    }
}

def t(key):
    lang = session.get('lang', 'pt')
    return TRANSLATIONS.get(lang, TRANSLATIONS['pt']).get(key, key)

@app.route('/set_language/<lang>')
def set_language(lang):
    if lang in ['en', 'pt']:
        session['lang'] = lang
    return redirect(request.referrer or url_for('index'))

def allowed_file(filename):
  return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def load_leave_requests():
  if os.path.exists(LEAVE_JSON_FILE):
    try:
      with open(LEAVE_JSON_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)
    except:
      return []
  return []

def save_leave_requests(leave_list):
  try:
    with open(LEAVE_JSON_FILE, 'w', encoding='utf-8') as f:
      json.dump(leave_list, f, indent=4, ensure_ascii=False)
  except Exception as e:
    print(f'Error saving {LEAVE_JSON_FILE}: {e}')

LEAVE_REQUESTS = load_leave_requests()

def get_emp_info(emp_code):
  emp_str = str(emp_code).strip()
  if not emp_str.startswith('NWC') and f'NWC{emp_str}' in MASTER_EMPLOYEES:
    emp_str = f'NWC{emp_str}'

  val = MASTER_EMPLOYEES.get(
      emp_str,
      {
          'name': f'Employee {emp_code}',
          'off': 'SUNDAY',
          'dept': 'General',
          'shift': 'morning',
      },
  )
  if isinstance(val, str):
    return {
        'name': val,
        'off': 'SUNDAY',
        'dept': 'General',
        'shift': 'morning',
    }
  return val

def upload_file_to_github(file_path, github_destination_path):
  token = os.getenv('GITHUB_TOKEN', '').strip()
  if not token:
    return False
  try:
    g = Github(token)
    repo = g.get_repo(GITHUB_REPO_NAME)
    with open(file_path, 'rb') as f:
      content = f.read()
    try:
      file_obj = repo.get_contents(github_destination_path, ref=GITHUB_BRANCH)
      repo.update_file(path=github_destination_path, message=f'Update file: {github_destination_path}', content=content, sha=file_obj.sha, branch=GITHUB_BRANCH)
    except Exception:
      repo.create_file(path=github_destination_path, message=f'Upload file: {github_destination_path}', content=content, branch=GITHUB_BRANCH)
    return True
  except Exception as e:
    print(f'[GITHUB ERROR]: {e}')
    return False

def save_leave_to_excel(emp_code, emp_name, start_date, end_date, leave_type, doc_filename):
  if os.path.exists(LEAVE_EXCEL_FILE):
    wb = openpyxl.load_workbook(LEAVE_EXCEL_FILE)
    ws = wb.active
  else:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Leave Records'
    ws.append(['Employee Code', 'Employee Name', 'Start Date', 'End Date', 'Leave Type', 'Document Name'])
  ws.append([emp_code, emp_name, start_date, end_date, leave_type, doc_filename])
  wb.save(LEAVE_EXCEL_FILE)

@app.route('/apply_leave', methods=['POST'])
def apply_leave():
  if not session.get('logged_in') or session.get('role') != 'employee':
    return redirect(url_for('login'))

  user_id = session.get('user_id')
  name = session.get('user_name')
  start_date = request.form.get('start_date')
  end_date = request.form.get('end_date')
  leave_type = request.form.get('leave_type', 'F10;1')

  filename = None
  file = request.files.get('supporting_doc')

  if file and file.filename != '':
    if not allowed_file(file.filename):
      flash('Invalid file format!', 'danger')
      return redirect(url_for('index'))
    filename = secure_filename(file.filename)
    local_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(local_path)
    upload_file_to_github(local_path, f'leave_documents/{user_id}_{filename}')

  save_leave_to_excel(user_id, name, start_date, end_date, leave_type, filename if filename else 'No Document')
  upload_file_to_github(LEAVE_EXCEL_FILE, 'leave_records.xlsx')

  leave_req = {
      'id': len(LEAVE_REQUESTS) + 1,
      'user_id': user_id,
      'name': name,
      'start_date': start_date,
      'end_date': end_date,
      'leave_type': leave_type,
      'filename': filename,
      'status': 'Pending',
  }
  LEAVE_REQUESTS.append(leave_req)
  save_leave_requests(LEAVE_REQUESTS)
  flash('Leave request successfully submitted!', 'success')
  return redirect(url_for('index'))

def check_device_connectivity():
  try:
    zk = ZK(MACHINE_IP, port=PORT, timeout=2, password=0, force_udp=False, ommit_ping=False)
    conn = zk.connect()
    if conn:
      conn.disconnect()
      return True
  except:
    pass
  if LAST_DEVICE_SYNC_TIME:
    if (datetime.now() - LAST_DEVICE_SYNC_TIME).total_seconds() < 300:
      return True
  return False

def fetch_attendance_data(start_date_str, end_date_str, filter_user_id):
  device_online = check_device_connectivity()
  period_data = {}
  raw_punches_list = []
  users_map_temp = {}

  for k, v in MASTER_EMPLOYEES.items():
    info = get_emp_info(k)
    code_formatted = f'NWC{k}' if not k.startswith('NWC') else k
    users_map_temp[str(k)] = {'code': code_formatted, 'name': info['name']}

  for uid_override, over_data in EMPLOYEE_OVERRIDES.items():
    users_map_temp[str(uid_override)] = {'code': over_data['code'], 'name': over_data['name']}

  attendance_records = []

  try:
    zk = ZK(MACHINE_IP, port=PORT, timeout=2, password=0, force_udp=False, ommit_ping=False)
    conn = zk.connect()
    if conn:
      users = conn.get_users()
      for user in users:
        uid_str = str(user.user_id)
        if uid_str in EMPLOYEE_OVERRIDES:
          emp_name = EMPLOYEE_OVERRIDES[uid_str]['name']
          emp_code = EMPLOYEE_OVERRIDES[uid_str]['code']
        else:
          info = get_emp_info(uid_str)
          emp_name = user.name if user.name else info['name']
          emp_code = f'NWC{uid_str}' if not uid_str.startswith('NWC') else uid_str
        users_map_temp[uid_str] = {'code': emp_code, 'name': emp_name}

      attendance = conn.get_attendance()
      for att in attendance:
        attendance_records.append({'user_id': str(att.user_id), 'timestamp': att.timestamp})
      conn.disconnect()
  except Exception as e:
    print(f'Device connection fallback: {e}')

  if not attendance_records and SYNCED_ATTENDANCE_LOGS:
    for log in SYNCED_ATTENDANCE_LOGS:
      ts = log['timestamp']
      if isinstance(ts, str):
        try:
          ts = datetime.strptime(ts, '%Y-%m-%d %H:%M:%S')
        except:
          continue
      attendance_records.append({'user_id': str(log['user_id']), 'timestamp': ts})

  for att in attendance_records:
    att_ts = att['timestamp']
    att_date_str = att_ts.strftime('%Y-%m-%d')
    if start_date_str <= att_date_str <= end_date_str:
      raw_uid = str(att['user_id'])
      if raw_uid in EMPLOYEE_OVERRIDES:
        emp_code = EMPLOYEE_OVERRIDES[raw_uid]['code']
        emp_name = EMPLOYEE_OVERRIDES[raw_uid]['name']
      elif raw_uid in users_map_temp:
        emp_code = users_map_temp[raw_uid]['code']
        emp_name = users_map_temp[raw_uid]['name']
      else:
        clean_uid = raw_uid.replace('NWC', '')
        info = get_emp_info(clean_uid)
        emp_name = info['name']
        emp_code = f'NWC{clean_uid}' if not clean_uid.startswith('NWC') else clean_uid

      if filter_user_id and filter_user_id != 'ALL' and emp_code != filter_user_id and raw_uid != filter_user_id:
        continue

      if att_date_str not in period_data:
        period_data[att_date_str] = {}
      if emp_code not in period_data[att_date_str]:
        period_data[att_date_str][emp_code] = {'name': emp_name, 'timestamps': []}
      
      if att_date_str in CUSTOM_PUNCHES and emp_code in CUSTOM_PUNCHES[att_date_str]:
        continue
      
      period_data[att_date_str][emp_code]['timestamps'].append(att_ts)

  for d_str, emp_dict in CUSTOM_PUNCHES.items():
    if start_date_str <= d_str <= end_date_str:
      for emp_code, times_list in emp_dict.items():
        if filter_user_id and filter_user_id != 'ALL' and emp_code != filter_user_id:
          continue
        if d_str not in period_data:
          period_data[d_str] = {}
        
        emp_info = get_emp_info(emp_code)
        parsed_times = []
        for t_str in times_list:
          try:
            dt_obj = datetime.strptime(f"{d_str} {t_str}", '%Y-%m-%d %H:%M:%S')
            parsed_times.append(dt_obj)
          except:
            pass
        parsed_times.sort()
        period_data[d_str][emp_code] = {'name': emp_info['name'], 'timestamps': parsed_times}

  for d_str, emp_dict in period_data.items():
    for emp_code, data_val in emp_dict.items():
      for ts_obj in data_val['timestamps']:
        raw_punches_list.append({
            'date': d_str,
            'time': ts_obj.strftime('%H:%M:%S'),
            'user_id': emp_code,
            'name': data_val['name'],
            'timestamp': ts_obj,
        })

  users_list = []
  for k, v in sorted(MASTER_EMPLOYEES.items(), key=lambda x: get_emp_info(x[0])['name']):
    info = get_emp_info(k)
    code_formatted = f'NWC{k}' if not k.startswith('NWC') else k
    users_list.append({
        'user_id': code_formatted,
        'name': info['name'],
        'dept': info['dept'],
        'off': info['off'],
    })

  final_data = []
  total_duration_seconds = 0
  total_lunch_seconds = 0
  total_net_variance_seconds = 0
  present_count, absent_count, off_count, mis_punch_count, late_arrival_count, ml_count = 0, 0, 0, 0, 0, 0
  shift_a_count, shift_b_count = 0, 0

  dates_to_process = sorted(period_data.keys(), reverse=True)
  if not dates_to_process and start_date_str == end_date_str:
    dates_to_process = [start_date_str]

  for date_str in dates_to_process:
    day_users_dict = period_data.get(date_str, {})
    current_dt = datetime.strptime(date_str, '%Y-%m-%d')
    current_day_name = current_dt.strftime('%A').upper()

    present_records, absent_records, off_records, mispunch_records, ml_records = [], [], [], [], []

    for emp_code, emp_data_val in MASTER_EMPLOYEES.items():
      final_emp_code = f'NWC{emp_code}' if not emp_code.startswith('NWC') else emp_code
      emp_info = get_emp_info(emp_code)
      emp_name, emp_dept, emp_shift = emp_info['name'], emp_info['dept'], emp_info['shift']
      
      emp_off = emp_info['off'].upper()
      if date_str in CUSTOM_OFFS and final_emp_code in CUSTOM_OFFS[date_str]:
        emp_off = CUSTOM_OFFS[date_str][final_emp_code].upper()

      if filter_user_id and filter_user_id != 'ALL' and final_emp_code != filter_user_id and emp_code != filter_user_id:
        continue

      approved_leave_obj = next((l for l in LEAVE_REQUESTS if l['user_id'] == final_emp_code and l['status'] == 'Approved' and l['start_date'] <= date_str <= l['end_date']), None)

      if approved_leave_obj:
        ml_count += 1
        leave_type_code = approved_leave_obj.get('leave_type', 'F10;1')
        ml_records.append({
            'date': date_str, 'user_id': final_emp_code, 'name': emp_name, 'dept': emp_dept,
            'store_in': f'Approved Leave ({leave_type_code})', 'lunch_out': '-', 'lunch_in': '-', 'out_time': '-',
            'total_lunch': '-', 'lunch_seconds': 3600, 'net_duration_seconds': 0, 'total_hours': '-',
            'net_variance': '-', 'variance_type': 'neutral', 'status': f'{leave_type_code} (Leave)',
            'is_late': 'No', 'shift_type': '-'
        })
        continue

      matched_key = emp_code if emp_code in day_users_dict else (final_emp_code if final_emp_code in day_users_dict else None)

      if matched_key:
        present_count += 1
        data_obj = day_users_dict[matched_key]
        sorted_times = sorted(data_obj['timestamps'])
        total_punches = len(sorted_times)

        store_in, lunch_out, lunch_in, out_time = '-', '-', '-', '-'
        lunch_seconds, net_duration_seconds = 0, 0
        total_lunch_str, total_hours_str, net_variance_str = '-', '-', '-'
        variance_type, status, is_late = 'neutral', 'Present', False
        shift_type = '-'

        first_punch_time = sorted_times[0].time()
        if first_punch_time <= time(10, 0, 0):
          shift_type = 'Shift A'
          shift_a_count += 1
        else:
          shift_type = 'Shift B'
          shift_b_count += 1

        if total_punches == 1:
          store_in = sorted_times[0].strftime('%H:%M:%S')
          status = 'Mis Punch'
          mis_punch_count += 1
        else:
          store_in = sorted_times[0].strftime('%H:%M:%S')
          out_time = sorted_times[-1].strftime('%H:%M:%S')
          store_in_time = sorted_times[0].time()
          limit_time = time(13, 10, 0) if emp_shift == 'second' else time(7, 0, 0)
          if store_in_time > limit_time:
            late_arrival_count += 1
            is_late = True

          if total_punches >= 3:
            lunch_out = sorted_times[1].strftime('%H:%M:%S')
            lunch_in = sorted_times[2].strftime('%H:%M:%S')
            actual_lunch_seconds = (sorted_times[2] - sorted_times[1]).seconds
            lunch_seconds = 3600 if actual_lunch_seconds < 3600 else actual_lunch_seconds
          else:
            lunch_seconds = 3600

          total_lunch_seconds += lunch_seconds
          l_hrs = divmod(lunch_seconds, 3600)
          total_lunch_str = f'{l_hrs[0]}h {l_hrs[1]//60}m'

          gross_seconds = (sorted_times[-1] - sorted_times[0]).seconds
          net_duration_seconds = max(0, gross_seconds - lunch_seconds)
          total_duration_seconds += net_duration_seconds

          hours = divmod(net_duration_seconds, 3600)
          total_hours_str = f'{hours[0]}h {hours[1]//60}m'

          target_seconds = 7 * 3600
          diff_from_target = net_duration_seconds - target_seconds
          total_net_variance_seconds += diff_from_target

          is_weekend = current_dt.weekday() >= 5
          if diff_from_target >= (45 * 60):
            e_hrs = divmod(diff_from_target, 3600)
            extra_hours_val = e_hrs[0] + (1 if e_hrs[1] >= 45*60 else 0)
            code_prefix = 'H07' if is_weekend else 'H06'
            net_variance_str, variance_type = f'{code_prefix};{extra_hours_val}', 'positive'
          elif diff_from_target > 0:
            net_variance_str, variance_type = '0h 0m', 'neutral'
          elif diff_from_target < 0:
            short_sec = abs(diff_from_target)
            s_hrs = divmod(short_sec, 3600)
            net_variance_str, variance_type = f'-{s_hrs[0]}h {s_hrs[1]//60}m', 'negative'
          else:
            net_variance_str, variance_type = '0h 0m', 'neutral'

          status = 'Present'

        if current_day_name == emp_off:
          status = 'Weekly Off'

        record = {
            'date': date_str, 'user_id': final_emp_code, 'name': emp_name, 'dept': emp_dept,
            'store_in': store_in, 'lunch_out': lunch_out, 'lunch_in': lunch_in, 'out_time': out_time,
            'total_lunch': total_lunch_str, 'lunch_seconds': lunch_seconds, 'net_duration_seconds': net_duration_seconds,
            'total_hours': total_hours_str, 'net_variance': net_variance_str, 'variance_type': variance_type,
            'status': status, 'is_late': 'Yes' if is_late else 'No', 'shift_type': shift_type,
        }

        if status == 'Weekly Off':
          off_records.append(record)
        elif status == 'Mis Punch':
          mispunch_records.append(record)
        else:
          present_records.append(record)
      else:
        if current_day_name == emp_off:
          off_count += 1
          off_records.append({
              'date': date_str, 'user_id': final_emp_code, 'name': emp_name, 'dept': emp_dept,
              'store_in': '-', 'lunch_out': '-', 'lunch_in': '-', 'out_time': '-', 'total_lunch': '-',
              'lunch_seconds': 3600, 'net_duration_seconds': 0, 'total_hours': '-', 'net_variance': 'Off',
              'variance_type': 'neutral', 'status': 'Weekly Off', 'is_late': 'No', 'shift_type': '-',
          })
        else:
          absent_count += 1
          absent_records.append({
              'date': date_str, 'user_id': final_emp_code, 'name': emp_name, 'dept': emp_dept,
              'store_in': '-', 'lunch_out': '-', 'lunch_in': '-', 'out_time': '-', 'total_lunch': '-',
              'lunch_seconds': 3600, 'net_duration_seconds': 0, 'total_hours': '-', 'net_variance': '-',
              'variance_type': 'neutral', 'status': 'Absent', 'is_late': 'No', 'shift_type': '-',
          })

    final_data.extend(present_records + mispunch_records + off_records + absent_records + ml_records)

  tot_hrs = divmod(total_duration_seconds, 3600)
  grand_total_hours = f'{tot_hrs[0]}h {tot_hrs[1]//60}m'
  tot_l_hrs = divmod(total_lunch_seconds, 3600)
  grand_total_lunch_hours = f'{tot_l_hrs[0]}h {tot_l_hrs[1]//60}m'

  v_sec = total_net_variance_seconds
  if v_sec >= 0:
    v_hrs = divmod(v_sec, 3600)
    grand_total_variance, grand_variance_type = f'+{v_hrs[0]}h {v_hrs[1]//60}m', 'positive'
  else:
    v_hrs = divmod(abs(v_sec), 3600)
    grand_total_variance, grand_variance_type = f'-{v_hrs[0]}h {v_hrs[1]//60}m', 'negative'

  stats_summary = {
      'present': present_count, 'absent': absent_count, 'off': off_count, 'mispunch': mis_punch_count,
      'late_arrival': late_arrival_count, 'ml': ml_count, 'shift_a': shift_a_count, 'shift_b': shift_b_count,
      'total_hrs': grand_total_hours, 'total_lunch_hrs': grand_total_lunch_hours, 'total_variance': grand_total_variance,
      'variance_type': grand_variance_type, 'device_online': device_online,
  }

  raw_punches_list = sorted(raw_punches_list, key=lambda x: x['timestamp'], reverse=True)
  return final_data, users_list, grand_total_hours, grand_total_lunch_hours, grand_total_variance, raw_punches_list, stats_summary


LOGIN_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Login | Gamek Attendance Portal</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style>body { font-family: 'Inter', sans-serif; }</style>
</head>
<body class="bg-slate-900 min-h-screen flex items-center justify-center p-4">
    <div class="bg-white/95 backdrop-blur-md rounded-3xl shadow-2xl border border-slate-200/50 p-8 w-full max-w-md space-y-6">
        <div class="flex justify-end space-x-2">
            <a href="/set_language/pt" class="px-2.5 py-1 rounded text-xs font-bold {% if session.get('lang', 'pt') == 'pt' %}bg-emerald-600 text-white{% else %}bg-slate-100 text-slate-700{% endif %}">PT</a>
            <a href="/set_language/en" class="px-2.5 py-1 rounded text-xs font-bold {% if session.get('lang', 'pt') == 'en' %}bg-emerald-600 text-white{% else %}bg-slate-100 text-slate-700{% endif %}">EN</a>
        </div>
        <div class="text-center space-y-2">
            <div class="inline-flex bg-[#78b13f] px-5 py-3 rounded-2xl shadow-lg mb-2 items-center justify-center">
                <img src="{{ url_for('static', filename='fresmart.png') }}" alt="Gamek Fresmart Logo" class="h-12 object-contain">
            </div>
            <h1 class="text-2xl font-black text-slate-900 tracking-tight">Gamek Fresmart Express</h1>
            <p class="text-xs text-slate-500 font-medium">Developed by Sonu Kumar <span class="text-emerald-600 font-semibold">(NCSA0608)</span></p>
        </div>
        {% with messages = get_flashed_messages(with_categories=true) %}
            {% if messages %}
                {% for category, message in messages %}
                <div class="{% if category == 'success' %}bg-emerald-50 border-emerald-200 text-emerald-800{% else %}bg-rose-50 border-rose-200 text-rose-700{% endif %} border text-xs font-semibold p-3.5 rounded-xl text-center">
                    {{ message }}
                </div>
                {% endfor %}
            {% endif %}
        {% endwith %}
        {% if error %}
        <div class="bg-rose-50 border border-rose-200 text-rose-700 text-xs font-semibold p-3.5 rounded-xl text-center">
            {{ error }}
        </div>
        {% endif %}
        <form method="POST" action="/login" class="space-y-4">
            <div>
                <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">Store/Employee Code</label>
                <input type="text" name="user_id" required placeholder="NWC1234" class="w-full bg-slate-50 border border-slate-300 rounded-xl px-4 py-3 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
            </div>
            <div>
                <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">Password</label>
                <input type="password" name="password" required placeholder="Enter Password" class="w-full bg-slate-50 border border-slate-300 rounded-xl px-4 py-3 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
            </div>
            <button type="submit" class="w-full bg-emerald-600 hover:bg-emerald-700 text-white font-bold text-xs uppercase tracking-widest py-3.5 rounded-xl shadow-lg transition duration-200">
                Secure Login 🚀
            </button>
        </form>
    </div>
</body>
</html>
"""

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Gamek HRM Dashboard</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style>
        body { font-family: 'Inter', sans-serif; }
        td.nowrap-cell { white-space: nowrap; }
        .excel-table th, .excel-table td { border: 1px solid #e2e8f0 !important; }
        th.sortable { cursor: pointer; user-select: none; transition: background-color 0.15s ease; }
        th.sortable:hover { background-color: #cbd5e1; }
        .stat-card { cursor: pointer; transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1); position: relative; overflow: hidden; }
        .stat-card:hover { transform: translateY(-2px); box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05); }
    </style>
    <script>
        let inactivityTimer;
        const INACTIVITY_LIMIT = 10 * 60 * 1000;
        function resetInactivityTimer() {
            clearTimeout(inactivityTimer);
            inactivityTimer = setTimeout(() => {
                alert("Session timeout ho gaya hai. Dobara login karein.");
                window.location.href = "/logout";
            }, INACTIVITY_LIMIT);
        }
        window.onload = function() {
            ['mousemove', 'keypress', 'click', 'scroll', 'touchstart'].forEach(eventName => {
                document.addEventListener(eventName, resetInactivityTimer, true);
            });
            resetInactivityTimer();
        };
        function updateLiveClock() {
            const now = new Date();
            const dateStr = now.toLocaleDateString('en-US', { weekday: 'short', year: 'numeric', month: 'short', day: 'numeric' });
            let h = String(now.getHours()).padStart(2, '0'), m = String(now.getMinutes()).padStart(2, '0'), s = String(now.getSeconds()).padStart(2, '0');
            const clockEl = document.getElementById('live-digital-clock');
            if (clockEl) clockEl.innerText = dateStr + ' | ' + h + ':' + m + ':' + s;
        }
        let sortDirections = {};
        function sortTable(columnIndex, isNumeric = false) {
            const table = document.getElementById("attendance-table");
            if (!table) return;
            const tbody = table.tBodies[0];
            const rows = Array.from(tbody.querySelectorAll("tr"));
            if (rows.length <= 1) return;
            let dir = sortDirections[columnIndex] || 'asc';
            sortDirections[columnIndex] = (dir === 'asc') ? 'desc' : 'asc';
            rows.sort((rowA, rowB) => {
                let cellA = rowA.cells[columnIndex].innerText.trim();
                let cellB = rowB.cells[columnIndex].innerText.trim();
                if (isNumeric) {
                    let valA = parseFloat(cellA.replace(/[^0-9.-]+/g,"")) || 0;
                    let valB = parseFloat(cellB.replace(/[^0-9.-]+/g,"")) || 0;
                    return (dir === 'asc') ? valA - valB : valB - valA;
                } else {
                    return (dir === 'asc') ? cellA.localeCompare(cellB) : cellB.localeCompare(cellA);
                }
            });
            tbody.innerHTML = "";
            rows.forEach((row, index) => {
                if (row.cells[0]) row.cells[0].innerText = index + 1;
                tbody.appendChild(row);
            });
        }
        function filterByStatus(statusVal) {
            let table = document.getElementById('attendance-table');
            if (!table) return;
            let trs = table.tBodies[0].getElementsByTagName('tr');
            for (let i = 0; i < trs.length; i++) {
                if (statusVal === 'ALL') {
                    trs[i].style.display = "";
                } else if (statusVal === 'Late Arrival') {
                    trs[i].style.display = (trs[i].getAttribute('data-late') === 'Yes') ? "" : "none";
                } else if (statusVal === 'Shift A' || statusVal === 'Shift B') {
                    trs[i].style.display = (trs[i].getAttribute('data-shift') === statusVal) ? "" : "none";
                } else if (statusVal === 'ML') {
                    let sc = trs[i].getElementsByTagName('td')[12];
                    if (sc) {
                        let text = sc.textContent || sc.innerText;
                        trs[i].style.display = (text.includes('F01;1') || text.includes('F05;1') || text.includes('F10;1') || text.includes('Leave')) ? "" : "none";
                    }
                } else {
                    let sc = trs[i].getElementsByTagName('td')[12];
                    if (sc) {
                        trs[i].style.display = (sc.textContent || sc.innerText).includes(statusVal) ? "" : "none";
                    }
                }
            }
        }
        function filterTableSearch() {
            let input = document.getElementById('table-search-input').value.toLowerCase();
            let table = document.getElementById('attendance-table');
            if (!table) return;
            let trs = table.tBodies[0].getElementsByTagName('tr');
            for (let i = 0; i < trs.length; i++) {
                let idCell = trs[i].getElementsByTagName('td')[2];
                let nameCell = trs[i].getElementsByTagName('td')[3];
                let deptCell = trs[i].getElementsByTagName('td')[4];
                if (idCell && nameCell && deptCell) {
                    let text = (idCell.textContent + ' ' + nameCell.textContent + ' ' + deptCell.textContent).toLowerCase();
                    trs[i].style.display = (text.indexOf(input) > -1) ? "" : "none";
                }
            }
        }
        let timeLeft = 180, timerInterval;
        function startTimer() {
            clearInterval(timerInterval);
            timeLeft = 180;
            timerInterval = setInterval(function() {
                if (timeLeft <= 0) { window.location.reload(); }
                else {
                    let m = Math.floor(timeLeft / 60), s = timeLeft % 60;
                    let timerEl = document.getElementById('countdown-timer');
                    if (timerEl) timerEl.innerText = m + ':' + (s < 10 ? '0' : '') + s;
                    timeLeft -= 1;
                }
            }, 1000);
        }
        function openExportModal() { document.getElementById('export-modal').classList.remove('hidden'); }
        function closeExportModal() { document.getElementById('export-modal').classList.add('hidden'); }
        function openRosterModal() { document.getElementById('roster-modal').classList.remove('hidden'); }
        function closeRosterModal() { document.getElementById('roster-modal').classList.add('hidden'); }
        function openCalendarModal() { document.getElementById('calendar-modal').classList.remove('hidden'); }
        function closeCalendarModal() { document.getElementById('calendar-modal').classList.add('hidden'); }
        function openLeaveModal() { document.getElementById('leave-modal').classList.remove('hidden'); }
        function closeLeaveModal() { document.getElementById('leave-modal').classList.add('hidden'); }
        function openWorkingHourModal() { document.getElementById('working-hour-modal').classList.remove('hidden'); }
        function closeWorkingHourModal() { document.getElementById('working-hour-modal').classList.add('hidden'); }
        function openDeveloperModal() { document.getElementById('developer-modal').classList.remove('hidden'); }
        function closeDeveloperModal() { document.getElementById('developer-modal').classList.add('hidden'); }

        function secureShutdown() {
            let pwd = prompt("Server band karne ke liye password enter karein:");
            if (pwd) window.location.href = "/shutdown?pwd=" + encodeURIComponent(pwd);
        }
        document.addEventListener('DOMContentLoaded', function() {
            startTimer();
            setInterval(updateLiveClock, 1000);
            updateLiveClock();
        });
    </script>
</head>
<body class="bg-slate-50 text-slate-800 antialiased flex h-screen overflow-hidden">
    <aside class="w-64 bg-white border-r border-slate-200 flex flex-col justify-between hidden lg:flex z-20">
        <div>
            <div class="p-5 flex items-center space-x-3 border-b border-slate-100">
                <div class="bg-[#78b13f] p-2 rounded-xl shadow-sm">
                    <img src="{{ url_for('static', filename='fresmart.png') }}" alt="Logo" class="h-6 object-contain">
                </div>
                <div>
                    <h2 class="text-sm font-bold text-slate-900 leading-tight">Gamek HRM</h2>
                    <p class="text-[10px] text-slate-400 font-medium">Fresmart Express LM11</p>
                </div>
            </div>
            <div class="p-4 space-y-1">
                <p class="text-[10px] font-bold uppercase tracking-wider text-slate-400 px-3 mb-2">Main Menu</p>
                <a href="/" class="flex items-center space-x-3 px-3 py-2.5 rounded-xl bg-slate-900 text-white font-semibold text-xs shadow-sm">
                    <span>📊</span><span>{{ t('dashboard') }}</span>
                </a>
                <a href="#" onclick="openWorkingHourModal(); return false;" class="flex items-center space-x-3 px-3 py-2.5 rounded-xl text-slate-600 hover:bg-slate-100 font-medium text-xs transition">
                    <span>⏱️</span><span>Employee Total Working Hour</span>
                </a>
                <a href="#" onclick="openCalendarModal(); return false;" class="flex items-center space-x-3 px-3 py-2.5 rounded-xl text-slate-600 hover:bg-slate-100 font-medium text-xs transition">
                    <span>📅</span><span>{{ t('calendar') }}</span>
                </a>
                <a href="#" onclick="openLeaveModal(); return false;" class="flex items-center justify-between px-3 py-2.5 rounded-xl text-slate-600 hover:bg-slate-100 font-medium text-xs transition">
                    <div class="flex items-center space-x-3"><span>🏖️️</span><span>{{ t('leave_mgmt') }}</span></div>
                    {% if role in ['admin', 'developer'] and pending_leaves_count > 0 %}
                    <span class="bg-rose-500 text-white text-[10px] font-bold px-2 py-0.5 rounded-full animate-pulse">{{ pending_leaves_count }}</span>
                    {% endif %}
                </a>
                {% if role == 'developer' %}
                <p class="text-[10px] font-bold uppercase tracking-wider text-rose-500 px-3 mt-6 mb-2">Developer Controls</p>
                <a href="#" onclick="openDeveloperModal(); return false;" class="flex items-center space-x-3 px-3 py-2.5 rounded-xl bg-rose-50 text-rose-700 font-bold text-xs transition border border-rose-200">
                    <span>🛠️</span><span>Developer Portal Panel</span>
                </a>
                {% endif %}
            </div>
        </div>
        <div class="p-4 border-t border-slate-100">
            <div class="bg-emerald-50 border border-emerald-100 rounded-2xl p-3.5 space-y-2">
                <div class="flex items-center space-x-2 text-emerald-800 font-bold text-xs"><span>📢</span><span>{{ t('announcements') }}</span></div>
                <p class="text-[11px] text-slate-600 leading-tight">{{ t('biometric_active') }}</p>
                <div class="text-[10px] text-emerald-600 font-bold pt-1">Dev: Sonu Kumar (NCSA0608)</div>
            </div>
        </div>
    </aside>

    <div class="flex-1 flex flex-col h-screen overflow-hidden">
        <header class="bg-white border-b border-slate-200 px-6 py-3.5 flex justify-between items-center z-10">
            <div class="flex items-center space-x-3">
                <h1 class="text-base font-black text-slate-900 tracking-tight">{{ t('dashboard') }}</h1>
                <span class="text-xs text-slate-400 font-medium">| {{ t('good_day') }}, {{ logged_user_name }}</span>
            </div>
            <div class="flex items-center space-x-3 flex-wrap">
                <div class="flex items-center bg-slate-100 p-1 rounded-xl border border-slate-200">
                    <a href="/set_language/pt" class="px-2 py-1 rounded-lg text-xs font-bold {% if session.get('lang', 'pt') == 'pt' %}bg-emerald-600 text-white shadow-sm{% else %}text-slate-600{% endif %}">PT</a>
                    <a href="/set_language/en" class="px-2 py-1 rounded-lg text-xs font-bold {% if session.get('lang', 'pt') == 'en' %}bg-emerald-600 text-white shadow-sm{% else %}text-slate-600{% endif %}">EN</a>
                </div>
                <button onclick="openWorkingHourModal()" class="bg-indigo-50 hover:bg-indigo-100 text-indigo-700 text-xs font-bold px-3 py-2 rounded-xl transition border border-indigo-200">
                    ⏱️ Working Hours
                </button>
                {% if role == 'developer' %}
                <button onclick="openDeveloperModal()" class="bg-rose-600 hover:bg-rose-700 text-white text-xs font-bold px-3 py-2 rounded-xl shadow transition animate-pulse">
                    🛠️ Developer Panel
                </button>
                {% endif %}
                <div class="text-xs bg-slate-50 px-3 py-2 rounded-xl border border-slate-200 flex items-center space-x-2">
                    <span class="h-2 w-2 {% if stats.device_online %}bg-emerald-500{% else %}bg-red-500{% endif %} rounded-full animate-pulse"></span>
                    <span class="text-slate-600 font-medium">Device: <strong class="{% if stats.device_online %}text-emerald-600{% else %}text-red-600{% endif %}">{% if stats.device_online %}{{ t('online') }}{% else %}{{ t('offline') }}{% endif %}</strong></span>
                    <span class="text-slate-300">|</span>
                    <span id="live-digital-clock" class="text-slate-700 font-semibold"></span>
                </div>
                <div class="flex items-center space-x-2 bg-slate-100 border border-slate-200 px-3 py-1.5 rounded-xl text-xs font-bold text-slate-700">
                    <span>👤 {{ logged_user_name }}</span>
                    <a href="/logout" class="text-rose-600 hover:text-rose-700 ml-2 font-semibold">{{ t('logout') }} 🔒</a>
                </div>
            </div>
        </header>

        <main class="flex-1 overflow-y-auto p-6 space-y-6">
            {% with messages = get_flashed_messages(with_categories=true) %}
                {% if messages %}
                    {% for category, message in messages %}
                    <div class="{% if category == 'success' %}bg-emerald-50 border-emerald-200 text-emerald-800{% else %}bg-rose-50 border-rose-200 text-rose-700{% endif %} border text-xs font-bold p-4 rounded-2xl shadow-sm flex items-center justify-between">
                        <span>{{ message }}</span><span class="cursor-pointer" onclick="this.parentElement.style.display='none'">✕</span>
                    </div>
                    {% endfor %}
                {% endif %}
            {% endwith %}

            <div class="grid grid-cols-2 sm:grid-cols-3 md:grid-cols-5 lg:grid-cols-9 gap-3">
                <div onclick="filterByStatus('Present')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-emerald-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('present') }}</p><h3 class="text-xl font-black text-emerald-600 mt-0.5">{{ stats.present }}</h3></div>
                    <div class="p-2 bg-emerald-50 text-emerald-600 rounded-xl">✅</div>
                </div>
                <div onclick="filterByStatus('Absent')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-rose-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('absent') }}</p><h3 class="text-xl font-black text-rose-600 mt-0.5">{{ stats.absent }}</h3></div>
                    <div class="p-2 bg-rose-50 text-rose-600 rounded-xl">❌</div>
                </div>
                <div onclick="filterByStatus('ML')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-cyan-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('medical_leave') }}</p><h3 class="text-xl font-black text-cyan-600 mt-0.5">{{ stats.ml }}</h3></div>
                    <div class="p-2 bg-cyan-50 text-cyan-600 rounded-xl">🏥</div>
                </div>
                <div onclick="filterByStatus('Weekly Off')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-slate-400">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('weekly_off') }}</p><h3 class="text-xl font-black text-slate-700 mt-0.5">{{ stats.off }}</h3></div>
                    <div class="p-2 bg-slate-100 text-slate-600 rounded-xl">🏖️</div>
                </div>
                <div onclick="filterByStatus('Late Arrival')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-amber-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('late_arrival') }}</p><h3 class="text-xl font-black text-amber-600 mt-0.5">{{ stats.late_arrival }}</h3></div>
                    <div class="p-2 bg-amber-50 text-amber-600 rounded-xl">⏰</div>
                </div>
                <div onclick="filterByStatus('Mis Punch')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-orange-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('mispunches') }}</p><h3 class="text-xl font-black text-orange-600 mt-0.5">{{ stats.mispunch }}</h3></div>
                    <div class="p-2 bg-orange-50 text-orange-600 rounded-xl">⚠</div>
                </div>
                <div onclick="filterByStatus('Shift A')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-blue-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">Shift A</p><h3 class="text-xl font-black text-blue-600 mt-0.5">{{ stats.shift_a }}</h3></div>
                    <div class="p-2 bg-blue-50 text-blue-600 rounded-xl">☀️</div>
                </div>
                <div onclick="filterByStatus('Shift B')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-indigo-500">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">Shift B</p><h3 class="text-xl font-black text-indigo-600 mt-0.5">{{ stats.shift_b }}</h3></div>
                    <div class="p-2 bg-indigo-50 text-indigo-600 rounded-xl">🌙</div>
                </div>
                <div onclick="filterByStatus('ALL')" class="stat-card bg-white p-3.5 rounded-2xl shadow-sm border border-slate-200 flex items-center justify-between border-l-4 border-l-emerald-600">
                    <div><p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">{{ t('total_hrs') }}</p><h3 class="text-xl font-black text-emerald-600 mt-0.5">{{ stats.total_hrs }}</h3></div>
                    <div class="p-2 bg-emerald-50 text-emerald-600 rounded-xl">⏱</div>
                </div>
            </div>

            <div class="bg-white rounded-2xl shadow-sm border border-slate-200 p-5">
                <form id="filter-form" method="GET" action="/" class="grid grid-cols-1 md:grid-cols-4 gap-4 items-end">
                    <div>
                        <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">{{ t('start_date') }}</label>
                        <input type="date" name="start_date" value="{{ start_date }}" onchange="this.form.submit()" class="w-full bg-slate-50 border border-slate-300 rounded-xl px-3.5 py-2.5 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
                    </div>
                    <div>
                        <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">{{ t('end_date') }}</label>
                        <input type="date" name="end_date" value="{{ end_date }}" onchange="this.form.submit()" class="w-full bg-slate-50 border border-slate-300 rounded-xl px-3.5 py-2.5 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
                    </div>
                    {% if role in ['admin', 'developer'] %}
                    <div>
                        <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">{{ t('emp_filter') }}</label>
                        <select name="employee" onchange="this.form.submit()" class="w-full bg-slate-50 border border-slate-300 rounded-xl px-3.5 py-2.5 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
                            <option value="ALL">{{ t('all_personnel') }}</option>
                            {% for emp in all_users %}
                                <option value="{{ emp.user_id }}" {% if selected_emp == emp.user_id %}selected{% endif %}>{{ emp.name }} ({{ emp.user_id }})</option>
                            {% endfor %}
                        </select>
                    </div>
                    {% else %}
                    <div>
                        <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">Logged In As</label>
                        <input type="hidden" name="employee" value="{{ selected_emp }}">
                        <input type="text" disabled value="{{ selected_emp }}" class="w-full bg-slate-100 border border-slate-200 rounded-xl px-3.5 py-2.5 text-sm font-bold text-emerald-700 cursor-not-allowed">
                    </div>
                    {% endif %}
                    <div class="flex space-x-2">
                        {% if role in ['admin', 'developer'] %}
                        <button type="button" onclick="openExportModal()" class="flex-1 bg-emerald-600 hover:bg-emerald-700 text-white font-bold text-xs uppercase tracking-wider py-3 px-3 rounded-xl text-center shadow-md transition">{{ t('export') }} 📥</button>
                        {% endif %}
                        <button type="button" onclick="openCalendarModal()" class="bg-slate-900 hover:bg-slate-800 text-white font-bold text-xs uppercase tracking-wider py-3 px-3 rounded-xl text-center shadow-md transition">📅 Rota</button>
                    </div>
                </form>
            </div>

            <div class="bg-white rounded-2xl shadow-sm border border-slate-200 overflow-hidden">
                <div class="p-4 bg-slate-50 border-b border-slate-200 flex flex-col sm:flex-row justify-between items-center gap-3">
                    <div class="text-xs text-slate-500 font-semibold">💡 {{ t('shift_hint') }}</div>
                    <div><input type="text" id="table-search-input" onkeyup="filterTableSearch()" placeholder="{{ t('search_placeholder') }}" class="bg-white border border-slate-300 rounded-xl px-3.5 py-2 text-xs w-72 shadow-sm font-medium focus:outline-none focus:border-emerald-500"></div>
                </div>
                <div class="overflow-x-auto">
                    <table id="attendance-table" class="w-full text-left border-collapse excel-table">
                        <thead>
                            <tr class="bg-slate-100 text-slate-600 uppercase text-[11px] font-bold tracking-wider">
                                <th class="py-3 px-4">{{ t('sr') }}</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(1)">{{ t('date') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(2)">{{ t('id') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(3)">{{ t('emp_name') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(4)">{{ t('dept') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(5)">{{ t('store_in') }} ↕</th>
                                <th class="py-3 px-4">{{ t('lunch_out') }}</th>
                                <th class="py-3 px-4">{{ t('lunch_in') }}</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(8)">{{ t('out_time') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(9)">{{ t('total_lunch') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(10)">{{ t('working_hours') }} ↕</th>
                                <th class="py-3 px-4 sortable" onclick="sortTable(11)">{{ t('overtime') }} ↕</th>
                                <th class="py-3 px-4 text-center sortable" onclick="sortTable(12)">{{ t('status') }} ↕</th>
                            </tr>
                        </thead>
                        <tbody class="text-sm text-slate-700 divide-y divide-slate-100">
                            {% if logs %}
                                {% for log in logs %}
                                <tr class="hover:bg-slate-50 transition-colors" data-late="{{ log.is_late }}" data-shift="{{ log.shift_type }}">
                                    <td class="py-3 px-4 nowrap-cell font-medium text-slate-400">{{ loop.index }}</td>
                                    <td class="py-3 px-4 nowrap-cell font-medium">{{ log.date }}</td>
                                    <td class="py-3 px-4 nowrap-cell text-slate-500 font-mono text-xs">{{ log.user_id }}</td>
                                    <td class="py-3 px-4 font-bold text-slate-900 nowrap-cell">{{ log.name }}</td>
                                    <td class="py-3 px-4 nowrap-cell"><span class="px-2.5 py-1 rounded-lg bg-emerald-50 text-emerald-700 text-[11px] font-bold">{{ log.dept }}</span></td>
                                    <td class="py-3 px-4 nowrap-cell font-mono text-xs">
                                        {{ log.store_in }}
                                        {% if log.shift_type == 'Shift A' %}<span class="text-[10px] bg-blue-100 text-blue-700 px-1.5 py-0.5 rounded font-bold ml-1">Shift A</span>
                                        {% elif log.shift_type == 'Shift B' %}<span class="text-[10px] bg-indigo-100 text-indigo-700 px-1.5 py-0.5 rounded font-bold ml-1">Shift B</span>{% endif %}
                                    </td>
                                    <td class="py-3 px-4 nowrap-cell font-mono text-xs text-slate-500">{{ log.lunch_out }}</td>
                                    <td class="py-3 px-4 nowrap-cell font-mono text-xs text-slate-500">{{ log.lunch_in }}</td>
                                    <td class="py-3 px-4 nowrap-cell font-mono text-xs">{{ log.out_time }}</td>
                                    <td class="py-3 px-4 font-bold nowrap-cell font-mono text-xs {% if log.lunch_seconds > 3600 %}text-rose-600 bg-rose-50/50{% else %}text-slate-700{% endif %}">{{ log.total_lunch }}</td>
                                    <td class="py-3 px-4 font-bold text-slate-900 nowrap-cell font-mono text-xs">{{ log.total_hours }}</td>
                                    <td class="py-3 px-4 font-bold nowrap-cell font-mono text-xs {% if log.variance_type == 'positive' %}text-emerald-600{% elif log.variance_type == 'negative' %}text-rose-600{% else %}text-slate-600{% endif %}">{{ log.net_variance }}</td>
                                    <td class="py-3 px-4 text-center nowrap-cell">
                                        {% if log.status == 'Weekly Off' %}<span class="px-3 py-1 rounded-full text-xs font-bold bg-slate-100 text-slate-600 border border-slate-200">Weekly Off</span>
                                        {% elif log.status == 'Absent' %}<span class="px-3 py-1 rounded-full text-xs font-bold bg-rose-50 text-rose-600 border border-rose-200">Absent</span>
                                        {% elif log.status == 'Present' %}<span class="px-3 py-1 rounded-full text-xs font-bold bg-emerald-50 text-emerald-700 border border-emerald-200">Present</span>
                                        {% else %}<span class="px-3 py-1 rounded-full text-xs font-bold bg-cyan-50 text-cyan-700 border border-cyan-200">{{ log.status }}</span>{% endif %}
                                    </td>
                                </tr>
                                {% endfor %}
                            {% else %}
                                <tr><td colspan="13" class="text-center py-16 text-slate-400 font-medium">No records found.</td></tr>
                            {% endif %}
                        </tbody>
                        <tfoot class="bg-slate-100 font-bold text-slate-900 text-sm border-t border-slate-200">
                            <tr>
                                <td colspan="9" class="py-4 px-4 text-right uppercase text-xs tracking-wider text-slate-500">{{ t('total_summary') }}</td>
                                <td class="py-4 px-4 text-emerald-700 font-mono">{{ grand_total_lunch_hours }}</td>
                                <td class="py-4 px-4 text-slate-900 font-mono">{{ grand_total_hours }}</td>
                                <td class="py-4 px-4 font-mono {% if stats.variance_type == 'positive' %}text-emerald-700{% else %}text-rose-700{% endif %}" colspan="2">{{ grand_total_variance }}</td>
                            </tr>
                        </tfoot>
                    </table>
                </div>
            </div>
        </main>
    </div>

    <div id="working-hour-modal" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center hidden">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 p-6 w-full max-w-xl mx-4 space-y-6">
            <div class="flex justify-between items-center border-b border-slate-100 pb-4">
                <h3 class="text-lg font-bold text-slate-900">⏱️ Employee Total Working Hour (Without Lunch)</h3>
                <button onclick="closeWorkingHourModal()" class="text-slate-400 hover:text-slate-600 font-bold text-lg">✕</button>
            </div>
            <form method="GET" action="/employee_working_hours" class="space-y-4">
                <div>
                    <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-2">Select Employee</label>
                    <select name="target_emp" required class="w-full bg-slate-50 border border-slate-300 rounded-xl px-4 py-3 text-sm font-medium focus:ring-2 focus:ring-emerald-500 focus:outline-none">
                        <option value="">-- Choose Employee --</option>
                        {% for emp in all_users %}
                            <option value="{{ emp.user_id }}">{{ emp.name }} ({{ emp.user_id }})</option>
                        {% endfor %}
                    </select>
                </div>
                <button type="submit" class="w-full bg-indigo-600 hover:bg-indigo-700 text-white font-bold text-xs uppercase tracking-widest py-3.5 rounded-xl shadow-lg transition">
                    View Working Hours Summary 🚀
                </button>
            </form>
            <div class="pt-2 flex justify-end">
                <button onclick="closeWorkingHourModal()" class="px-4 py-2 bg-slate-100 hover:bg-slate-200 text-slate-700 text-xs font-bold rounded-xl transition">Close</button>
            </div>
        </div>
    </div>

    {% if role == 'developer' %}
    <div id="developer-modal" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center hidden">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 p-6 w-full max-w-3xl mx-4 space-y-6 max-h-[90vh] overflow-y-auto">
            <div class="flex justify-between items-center border-b border-slate-100 pb-4">
                <h3 class="text-lg font-bold text-rose-700 flex items-center gap-2">🛠️ Developer Portal Control Panel</h3>
                <button onclick="closeDeveloperModal()" class="text-slate-400 hover:text-slate-600 font-bold text-lg">✕</button>
            </div>
            
            <div class="space-y-6">
                <div class="bg-slate-50 border border-slate-200 rounded-2xl p-5 space-y-4">
                    <h4 class="text-sm font-bold text-slate-900">1. Edit Employee Punching Hours (Reflects Everywhere)</h4>
                    <form method="POST" action="/developer_update_punch" class="grid grid-cols-1 sm:grid-cols-2 gap-4">
                        <div>
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">Select Employee</label>
                            <select name="emp_code" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-medium focus:ring-2 focus:ring-rose-500">
                                {% for emp in all_users %}
                                <option value="{{ emp.user_id }}">{{ emp.name }} ({{ emp.user_id }})</option>
                                {% endfor %}
                            </select>
                        </div>
                        <div>
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">Date</label>
                            <input type="date" name="punch_date" required value="{{ datetime.now().strftime('%Y-%m-%d') }}" class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-medium">
                        </div>
                        <div class="sm:col-span-2">
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">Punches (Comma separated HH:MM:SS format)</label>
                            <input type="text" name="punches_str" required placeholder="08:00:00, 12:00:00, 13:00:00, 17:00:00" class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-mono">
                            <p class="text-[11px] text-slate-400 mt-1">Note: 2 punches = Store In & Out, 4 punches = Store In, Lunch Out, Lunch In, Out Time.</p>
                        </div>
                        <div class="sm:col-span-2 flex justify-end">
                            <button type="submit" class="bg-rose-600 hover:bg-rose-700 text-white font-bold text-xs uppercase tracking-wider px-5 py-2.5 rounded-xl shadow transition">Submit & Update Punches 💾</button>
                        </div>
                    </form>
                </div>

                <div class="bg-slate-50 border border-slate-200 rounded-2xl p-5 space-y-4">
                    <h4 class="text-sm font-bold text-slate-900">2. Change Date-wise Weekly Off</h4>
                    <form method="POST" action="/developer_update_off" class="grid grid-cols-1 sm:grid-cols-3 gap-4">
                        <div>
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">Select Employee</label>
                            <select name="emp_code" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-medium">
                                {% for emp in all_users %}
                                <option value="{{ emp.user_id }}">{{ emp.name }} ({{ emp.user_id }})</option>
                                {% endfor %}
                            </select>
                        </div>
                        <div>
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">Specific Date</label>
                            <input type="date" name="off_date" required value="{{ datetime.now().strftime('%Y-%m-%d') }}" class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-medium">
                        </div>
                        <div>
                            <label class="block text-xs font-bold uppercase tracking-wider text-slate-500 mb-1">New Weekly Off Day</label>
                            <select name="new_off" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs font-medium">
                                <option value="SUNDAY">SUNDAY</option>
                                <option value="MONDAY">MONDAY</option>
                                <option value="TUESDAY">TUESDAY</option>
                                <option value="WEDNESDAY">WEDNESDAY</option>
                                <option value="THURSDAY">THURSDAY</option>
                                <option value="FRIDAY">FRIDAY</option>
                                <option value="SATURDAY">SATURDAY</option>
                            </select>
                        </div>
                        <div class="sm:col-span-3 flex justify-end">
                            <button type="submit" class="bg-indigo-600 hover:bg-indigo-700 text-white font-bold text-xs uppercase tracking-wider px-5 py-2.5 rounded-xl shadow transition">Update Weekly Off for Date 📅</button>
                        </div>
                    </form>
                </div>
            </div>

            <div class="pt-2 flex justify-end">
                <button onclick="closeDeveloperModal()" class="px-4 py-2 bg-slate-100 hover:bg-slate-200 text-slate-700 text-xs font-bold rounded-xl transition">Close</button>
            </div>
        </div>
    </div>
    {% endif %}

    <div id="leave-modal" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center hidden">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 p-6 w-full max-w-4xl mx-4 space-y-6 max-h-[85vh] flex flex-col">
            <div class="flex justify-between items-center border-b border-slate-100 pb-4">
                <h3 class="text-lg font-bold text-slate-900 flex items-center gap-2">🏖️ {{ t('leave_mgmt') }}</h3>
                <button onclick="closeLeaveModal()" class="text-slate-400 hover:text-slate-600 font-bold text-lg">✕</button>
            </div>
            <div class="overflow-y-auto flex-1 space-y-6">
                {% if role == 'employee' %}
                <div class="bg-slate-50 border border-slate-200 rounded-2xl p-5 space-y-4">
                    <h4 class="text-sm font-bold text-slate-900">{{ t('apply_leave') }}</h4>
                    <form method="POST" action="/apply_leave" enctype="multipart/form-data" class="grid grid-cols-1 sm:grid-cols-3 gap-4">
                        <div><label class="block text-xs font-bold uppercase text-slate-500 mb-1">{{ t('start_date') }}</label><input type="date" name="start_date" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs"></div>
                        <div><label class="block text-xs font-bold uppercase text-slate-500 mb-1">{{ t('end_date') }}</label><input type="date" name="end_date" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs"></div>
                        <div>
                            <label class="block text-xs font-bold uppercase text-slate-500 mb-1">{{ t('leave_type') }}</label>
                            <select name="leave_type" required class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs">
                                <option value="F01;1">F01;1 - Baixa Médica</option>
                                <option value="F03;1">F03;1 - Falta Injustificada</option>
                                <option value="F05;1">F05;1 - Licença sem vencimento</option>
                                <option value="F10;1" selected>F10;1 - Falta Justificada</option>
                            </select>
                        </div>
                        <div class="sm:col-span-3"><label class="block text-xs font-bold uppercase text-slate-500 mb-1">{{ t('support_doc') }}</label><input type="file" name="supporting_doc" accept=".pdf,image/*" class="w-full bg-white border border-slate-300 rounded-xl px-3 py-2 text-xs"></div>
                        <div class="sm:col-span-3 flex justify-end"><button type="submit" class="bg-emerald-600 hover:bg-emerald-700 text-white font-bold text-xs uppercase tracking-wider px-5 py-2.5 rounded-xl shadow transition">{{ t('submit_req') }} 🚀</button></div>
                    </form>
                </div>
                {% endif %}
                <div class="space-y-3">
                    <h4 class="text-sm font-bold text-slate-900">{{ t('leave_history') }}</h4>
                    <div class="border border-slate-200 rounded-xl overflow-hidden">
                        <table class="w-full text-left border-collapse">
                            <thead>
                                <tr class="bg-slate-100 text-slate-600 uppercase text-[10px] font-bold tracking-wider">
                                    <th class="py-2.5 px-3 border-b">{{ t('emp_name') }}</th>
                                    <th class="py-2.5 px-3 border-b">Dates</th>
                                    <th class="py-2.5 px-3 border-b">Type</th>
                                    <th class="py-2.5 px-3 border-b text-center">{{ t('status') }}</th>
                                    {% if role in ['admin', 'developer'] %}<th class="py-2.5 px-3 border-b text-center">{{ t('action') }}</th>{% endif %}
                                </tr>
                            </thead>
                            <tbody class="text-xs text-slate-700 divide-y divide-slate-100">
                                {% if leave_requests %}
                                    {% for req in leave_requests %}
                                    <tr class="hover:bg-slate-50">
                                        <td class="py-2.5 px-3"><div class="font-bold text-slate-900">{{ req.name }}</div><div class="text-[10px] text-slate-400 font-mono">{{ req.user_id }}</div></td>
                                        <td class="py-2.5 px-3 font-mono text-[11px]">{{ req.start_date }} to {{ req.end_date }}</td>
                                        <td class="py-2.5 px-3 font-bold text-cyan-700">{{ req.leave_type }}</td>
                                        <td class="py-2.5 px-3 text-center"><span class="px-2.5 py-1 rounded-full {% if req.status == 'Approved' %}bg-emerald-50 text-emerald-700{% elif req.status == 'Rejected' %}bg-rose-50 text-rose-700{% else %}bg-amber-50 text-amber-700{% endif %} font-bold text-[10px]">{{ req.status }}</span></td>
                                        {% if role in ['admin', 'developer'] %}
                                        <td class="py-2.5 px-3 text-center">
                                            {% if req.status == 'Pending' %}
                                            <div class="flex items-center justify-center space-x-1.5">
                                                <a href="/update_leave/{{ req.id }}/approve" class="bg-emerald-600 text-white font-bold px-2.5 py-1 rounded-lg text-[10px]">{{ t('approve') }}</a>
                                                <a href="/update_leave/{{ req.id }}/reject" class="bg-rose-600 text-white font-bold px-2.5 py-1 rounded-lg text-[10px]">{{ t('reject') }}</a>
                                            </div>
                                            {% else %}<span class="text-slate-400 text-[10px]">-</span>{% endif %}
                                        </td>
                                        {% endif %}
                                    </tr>
                                    {% endfor %}
                                {% else %}<tr><td colspan="5" class="text-center py-8 text-slate-400">No records found.</td></tr>{% endif %}
                            </tbody>
                        </table>
                    </div>
                </div>
            </div>
            <div class="pt-2 flex justify-end"><button onclick="closeLeaveModal()" class="px-4 py-2 bg-slate-100 hover:bg-slate-200 text-slate-700 text-xs font-bold rounded-xl">{{ t('close') }}</button></div>
        </div>
    </div>

    <div id="export-modal" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center hidden">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 p-6 w-full max-w-md mx-4 space-y-6">
            <div class="flex justify-between items-center border-b border-slate-100 pb-4"><h3 class="text-lg font-bold text-slate-900">📊 Export Options</h3><button onclick="closeExportModal()" class="text-slate-400 font-bold text-lg">✕</button></div>
            <div class="space-y-4">
                <a href="/export?start_date={{ start_date }}&end_date={{ end_date }}&employee={{ selected_emp }}" onclick="closeExportModal()" class="block p-4 rounded-xl border border-slate-200 hover:border-emerald-500 hover:bg-emerald-50/50 transition font-bold text-slate-900">📄 Standard Row Export</a>
                <a href="/export_matrix?start_date={{ start_date }}&end_date={{ end_date }}&employee={{ selected_emp }}" onclick="closeExportModal()" class="block p-4 rounded-xl border border-slate-200 hover:border-emerald-500 hover:bg-emerald-50/50 transition font-bold text-slate-900">📅 Employee Matrix (Leave Codes)</a>
            </div>
            <div class="pt-2 flex justify-end"><button onclick="closeExportModal()" class="px-4 py-2 bg-slate-100 text-slate-700 text-xs font-bold rounded-xl">Cancel</button></div>
        </div>
    </div>

    <div id="calendar-modal" class="fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center hidden">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 p-6 w-full max-w-5xl mx-4 space-y-6 max-h-[85vh] flex flex-col">
            <div class="flex justify-between items-center border-b border-slate-100 pb-4"><h3 class="text-lg font-bold text-slate-900">📅 Calendar & Rota</h3><button onclick="closeCalendarModal()" class="text-slate-400 font-bold text-lg">✕</button></div>
            <div class="overflow-y-auto flex-1">
                <table class="w-full text-left border-collapse">
                    <thead>
                        <tr class="bg-slate-100 text-slate-600 uppercase text-[11px] font-bold tracking-wider">
                            <th class="py-3 px-4 border">Sr.</th><th class="py-3 px-4 border">ID</th><th class="py-3 px-4 border">Employee Name</th><th class="py-3 px-4 border">Department</th><th class="py-3 px-4 text-center border bg-indigo-50 text-indigo-800">Weekly Off Day</th>
                        </tr>
                    </thead>
                    <tbody class="text-xs text-slate-700 divide-y divide-slate-100">
                        {% for emp in all_users %}
                        <tr class="hover:bg-slate-50">
                            <td class="py-2.5 px-4 border text-slate-400">{{ loop.index }}</td>
                            <td class="py-2.5 px-4 border font-mono font-semibold">{{ emp.user_id }}</td>
                            <td class="py-2.5 px-4 border font-bold text-slate-900">{{ emp.name }}</td>
                            <td class="py-2.5 px-4 border"><span class="px-2 py-0.5 rounded bg-emerald-50 text-emerald-700 font-bold">{{ emp.dept }}</span></td>
                            <td class="py-2.5 px-4 border text-center font-bold text-indigo-700 bg-indigo-50/50 text-sm">{{ emp.off }}</td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
            <div class="pt-2 flex justify-end"><button onclick="closeCalendarModal()" class="px-4 py-2 bg-slate-100 text-slate-700 text-xs font-bold rounded-xl">Close</button></div>
        </div>
    </div>
</body>
</html>
"""

WORKING_HOURS_SUMMARY_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Employee Working Hours | Gamek HRM</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style>body { font-family: 'Inter', sans-serif; }</style>
</head>
<body class="bg-slate-900 min-h-screen flex items-center justify-center p-4">
    <div class="bg-white rounded-3xl shadow-2xl border border-slate-200/50 p-8 w-full max-w-2xl space-y-6">
        <div class="flex justify-between items-center border-b border-slate-100 pb-4">
            <div>
                <h1 class="text-xl font-black text-slate-900">{{ emp_name }} <span class="text-xs font-mono text-indigo-600">({{ emp_code }})</span></h1>
                <p class="text-xs text-slate-500 font-medium">Working Hours Summary (Without Lunch)</p>
            </div>
            <a href="/" class="bg-slate-100 hover:bg-slate-200 text-slate-700 font-bold text-xs px-4 py-2 rounded-xl transition">Back to Dashboard ⬅</a>
        </div>

        <div class="grid grid-cols-1 sm:grid-cols-2 gap-4">
            <div class="bg-slate-50 border border-slate-200 p-5 rounded-2xl flex items-center justify-between">
                <div>
                    <p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">Today</p>
                    <h3 class="text-2xl font-black text-indigo-600 mt-1">{{ wh_data.today }}</h3>
                </div>
                <div class="p-3 bg-indigo-50 text-indigo-600 rounded-xl text-lg">📅</div>
            </div>
            <div class="bg-slate-50 border border-slate-200 p-5 rounded-2xl flex items-center justify-between">
                <div>
                    <p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">Yesterday</p>
                    <h3 class="text-2xl font-black text-slate-700 mt-1">{{ wh_data.yesterday }}</h3>
                </div>
                <div class="p-3 bg-slate-100 text-slate-600 rounded-xl text-lg">⏳</div>
            </div>
            <div class="bg-slate-50 border border-slate-200 p-5 rounded-2xl flex items-center justify-between">
                <div>
                    <p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">This Week</p>
                    <h3 class="text-2xl font-black text-emerald-600 mt-1">{{ wh_data.this_week }}</h3>
                </div>
                <div class="p-3 bg-emerald-50 text-emerald-600 rounded-xl text-lg">📊</div>
            </div>
            <div class="bg-slate-50 border border-slate-200 p-5 rounded-2xl flex items-center justify-between">
                <div>
                    <p class="text-[10px] font-bold uppercase tracking-wider text-slate-400">Last Week</p>
                    <h3 class="text-2xl font-black text-blue-600 mt-1">{{ wh_data.last_week }}</h3>
                </div>
                <div class="p-3 bg-blue-50 text-blue-600 rounded-xl text-lg">📈</div>
            </div>
            <div class="sm:col-span-2 bg-emerald-50 border border-emerald-200 p-5 rounded-2xl flex items-center justify-between">
                <div>
                    <p class="text-[10px] font-bold uppercase tracking-wider text-emerald-800">This Month</p>
                    <h3 class="text-3xl font-black text-emerald-700 mt-1">{{ wh_data.this_month }}</h3>
                </div>
                <div class="p-3 bg-emerald-100 text-emerald-800 rounded-xl text-xl">🏆</div>
            </div>
        </div>
    </div>
</body>
</html>
"""

@app.route('/employee_working_hours')
def employee_working_hours():
  if not session.get('logged_in'):
    return redirect(url_for('login'))
  
  target_emp = request.args.get('target_emp')
  if not target_emp:
    flash('Please select an employee.', 'danger')
    return redirect(url_for('index'))

  emp_info = get_emp_info(target_emp)
  emp_name = emp_info['name']

  today_dt = datetime.now()
  today_str = today_dt.strftime('%Y-%m-%d')
  yesterday_str = (today_dt - timedelta(days=1)).strftime('%Y-%m-%d')
  start_of_this_week = (today_dt - timedelta(days=today_dt.weekday())).strftime('%Y-%m-%d')
  start_of_last_week = (today_dt - timedelta(days=today_dt.weekday() + 7)).strftime('%Y-%m-%d')
  end_of_last_week = (today_dt - timedelta(days=today_dt.weekday() + 1)).strftime('%Y-%m-%d')
  start_of_this_month = today_dt.replace(day=1).strftime('%Y-%m-%d')

  def calculate_net_seconds_for_range(start_s, end_s):
    logs, _, _, _, _, _, _ = fetch_attendance_data(start_s, end_s, target_emp)
    total_sec = 0
    for l in logs:
      if 'net_duration_seconds' in l and isinstance(l['net_duration_seconds'], (int, float)):
        total_sec += l['net_duration_seconds']
    h = divmod(total_sec, 3600)
    return f"{h[0]}h {h[1]//60}m"

  wh_data = {
      'today': calculate_net_seconds_for_range(today_str, today_str),
      'yesterday': calculate_net_seconds_for_range(yesterday_str, yesterday_str),
      'this_week': calculate_net_seconds_for_range(start_of_this_week, today_str),
      'last_week': calculate_net_seconds_for_range(start_of_last_week, end_of_last_week),
      'this_month': calculate_net_seconds_for_range(start_of_this_month, today_str),
  }

  return render_template_string(
      WORKING_HOURS_SUMMARY_TEMPLATE,
      emp_name=emp_name,
      emp_code=target_emp,
      wh_data=wh_data
  )

@app.route('/developer_update_punch', methods=['POST'])
def developer_update_punch():
  if not session.get('logged_in') or session.get('role') != 'developer':
    return redirect(url_for('login'))

  emp_code = request.form.get('emp_code')
  punch_date = request.form.get('punch_date')
  punches_str = request.form.get('punches_str', '')

  if not emp_code or not punch_date:
    flash('Employee aur Date zaroori hain.', 'danger')
    return redirect(url_for('index'))

  punches_list = [p.strip() for p in punches_str.split(',') if p.strip()]

  if punch_date not in CUSTOM_PUNCHES:
    CUSTOM_PUNCHES[punch_date] = {}
  CUSTOM_PUNCHES[punch_date][emp_code] = punches_list
  save_json_file(CUSTOM_PUNCHES_FILE, CUSTOM_PUNCHES)

  flash(f'Punching hours for {emp_code} on {punch_date} successfully updated and reflected for everyone!', 'success')
  return redirect(url_for('index'))

@app.route('/developer_update_off', methods=['POST'])
def developer_update_off():
  if not session.get('logged_in') or session.get('role') != 'developer':
    return redirect(url_for('login'))

  emp_code = request.form.get('emp_code')
  off_date = request.form.get('off_date')
  new_off = request.form.get('new_off')

  if not emp_code or not off_date or not new_off:
    flash('Sabhi fields bharna zaroori hain.', 'danger')
    return redirect(url_for('index'))

  if off_date not in CUSTOM_OFFS:
    CUSTOM_OFFS[off_date] = {}
  CUSTOM_OFFS[off_date][emp_code] = new_off
  save_json_file(CUSTOM_OFFS_FILE, CUSTOM_OFFS)

  flash(f'Weekly off for {emp_code} on {off_date} changed to {new_off} successfully!', 'success')
  return redirect(url_for('index'))

@app.route('/login', methods=['GET', 'POST'])
def login():
  if request.method == 'POST':
    uid = request.form.get('user_id').strip().upper()
    pwd = request.form.get('password').strip()

    admin_pass = os.getenv('ADMIN_PWD', 'Gamek@789')
    dev_pass = os.getenv('DEV_PWD', 'Shama@8577')

    if uid == 'LM11' and pwd == admin_pass:
      session['logged_in'] = True
      session['role'] = 'admin'
      session['user_id'] = 'LM11'
      session['user_name'] = 'Admin (LM11)'
      session.setdefault('lang', 'pt')
      return redirect(url_for('index'))

    if uid == 'NCSA0608' and pwd == dev_pass:
      session['logged_in'] = True
      session['role'] = 'developer'
      session['user_id'] = 'NCSA0608'
      session['user_name'] = 'Sonu Kumar (Developer)'
      session.setdefault('lang', 'pt')
      return redirect(url_for('index'))

    emp_key = f'NWC{uid}' if not uid.startswith('NWC') else uid
    if emp_key in MASTER_EMPLOYEES and pwd == '123':
      session['logged_in'] = True
      session['role'] = 'employee'
      session['user_id'] = emp_key
      session['user_name'] = MASTER_EMPLOYEES[emp_key]['name']
      session.setdefault('lang', 'pt')
      return redirect(url_for('index'))
    else:
      return render_template_string(LOGIN_TEMPLATE, error='Galat User ID ya Password!')
  return render_template_string(LOGIN_TEMPLATE, error=None)

@app.route('/logout')
def logout():
  role = session.get('role')
  user_name = session.get('user_name', '')
  msg_text = 'Thank you admin' if role == 'admin' else f'Thank you {user_name}'
  session.clear()
  flash(msg_text, 'success')
  return redirect(url_for('login'))

@app.route('/')
def index():
  if not session.get('logged_in'):
    return redirect(url_for('login'))

  role = session.get('role')
  logged_user_id = session.get('user_id')
  logged_user_name = session.get('user_name')

  today_str = datetime.now().strftime('%Y-%m-%d')
  start_date = request.args.get('start_date', today_str)
  end_date = request.args.get('end_date', today_str)

  if role == 'employee':
    selected_emp = logged_user_id
  else:
    selected_emp = request.args.get('employee', 'ALL')

  logs, users_list, grand_total_hours, grand_total_lunch_hours, grand_total_variance, raw_punches_list, stats = fetch_attendance_data(
      start_date, end_date, selected_emp
  )

  pending_leaves_count = sum(1 for l in LEAVE_REQUESTS if l['status'] == 'Pending')

  return render_template_string(
      HTML_TEMPLATE,
      logs=logs,
      all_users=users_list,
      grand_total_hours=grand_total_hours,
      grand_total_lunch_hours=grand_total_lunch_hours,
      grand_total_variance=grand_total_variance,
      stats=stats,
      start_date=start_date,
      end_date=end_date,
      selected_emp=selected_emp,
      role=role,
      logged_user_name=logged_user_name,
      pending_leaves_count=pending_leaves_count,
      leave_requests=LEAVE_REQUESTS,
      raw_punches=raw_punches_list
  )

@app.route('/update_leave/<int:leave_id>/<action>')
def update_leave(leave_id, action):
  if not session.get('logged_in') or session.get('role') not in ['admin', 'developer']:
    return redirect(url_for('login'))

  req = next((l for l in LEAVE_REQUESTS if l['id'] == leave_id), None)
  if req:
    if action == 'approve':
      req['status'] = 'Approved'
      flash('Leave request approved successfully!', 'success')
    elif action == 'reject':
      req['status'] = 'Rejected'
      flash('Leave request rejected.', 'danger')
    save_leave_requests(LEAVE_REQUESTS)
  return redirect(url_for('index'))

@app.route('/export')
def export():
  if not session.get('logged_in') or session.get('role') not in ['admin', 'developer']:
    return redirect(url_for('login'))

  start_date = request.args.get('start_date', datetime.now().strftime('%Y-%m-%d'))
  end_date = request.args.get('end_date', datetime.now().strftime('%Y-%m-%d'))
  employee = request.args.get('employee', 'ALL')

  logs, _, _, _, _, _, _ = fetch_attendance_data(start_date, end_date, employee)

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Attendance Report'

  header_font = Font(name='Inter', size=11, bold=True, color='FFFFFF')
  header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
  align_center = Alignment(horizontal='center', vertical='center')
  align_left = Alignment(horizontal='left', vertical='center')
  thin_border = Border(
      left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
      top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
  )

  headers = [
      'Sr. No.', 'Date', 'Employee ID', 'Employee Name', 'Department',
      'Store In', 'Lunch Out', 'Lunch In', 'Out Time', 'Total Lunch',
      'Working Hours', 'Overtime', 'Status'
  ]
  ws.append(headers)

  for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_center
    cell.border = thin_border

  for idx, log in enumerate(logs, start=2):
    row_data = [
        idx - 1, log['date'], log['user_id'], log['name'], log['dept'],
        log['store_in'], log['lunch_out'], log['lunch_in'], log['out_time'],
        log['total_lunch'], log['total_hours'], log['net_variance'], log['status']
    ]
    ws.append(row_data)
    for col_num in range(1, len(row_data) + 1):
      cell = ws.cell(row=idx, column=col_num)
      cell.border = thin_border
      cell.alignment = align_center if col_num in [1, 2, 3, 6, 7, 8, 9, 10, 11, 12, 13] else align_left

  export_filename = f'attendance_report_{start_date}_to_{end_date}.xlsx'
  export_path = os.path.join(app.config['UPLOAD_FOLDER'], export_filename)
  wb.save(export_path)
  return send_file(export_path, as_attachment=True)

@app.route('/export_matrix')
def export_matrix():
  if not session.get('logged_in') or session.get('role') not in ['admin', 'developer']:
    return redirect(url_for('login'))

  start_date_str = request.args.get('start_date', datetime.now().strftime('%Y-%m-%d'))
  end_date_str = request.args.get('end_date', datetime.now().strftime('%Y-%m-%d'))
  employee = request.args.get('employee', 'ALL')

  start_dt = datetime.strptime(start_date_str, '%Y-%m-%d')
  end_dt = datetime.strptime(end_date_str, '%Y-%m-%d')
  delta_days = (end_dt - start_dt).days

  dates_list = [(start_dt + timedelta(days=i)).strftime('%Y-%m-%d') for i in range(delta_days + 1)]

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Employee Matrix'

  header_font = Font(name='Inter', size=10, bold=True, color='FFFFFF')
  header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
  align_center = Alignment(horizontal='center', vertical='center')
  thin_border = Border(
      left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
      top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
  )

  headers = ['Emp ID', 'Employee Name', 'Department'] + dates_list
  ws.append(headers)

  for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_center
    cell.border = thin_border

  users_data = MASTER_EMPLOYEES.items()
  row_idx = 2
  for k, v in sorted(users_data, key=lambda x: get_emp_info(x[0])['name']):
    emp_code = f'NWC{k}' if not k.startswith('NWC') else k
    if employee != 'ALL' and employee != emp_code:
      continue
    info = get_emp_info(k)
    row_values = [emp_code, info['name'], info['dept']]
    for d_str in dates_list:
      logs, _, _, _, _, _, _ = fetch_attendance_data(d_str, d_str, emp_code)
      if logs:
        st = logs[0]['status']
        if st == 'Present':
          row_values.append('P')
        elif st == 'Absent':
          row_values.append('A')
        elif st == 'Weekly Off':
          row_values.append('OFF')
        elif st == 'Mis Punch':
          row_values.append('MP')
        else:
          row_values.append(st[:10])
      else:
        row_values.append('-')
    ws.append(row_values)
    for col_num in range(1, len(row_values) + 1):
      cell = ws.cell(row=row_idx, column=col_num)
      cell.border = thin_border
      cell.alignment = align_center if col_num > 3 else Alignment(horizontal='left', vertical='center')
    row_idx += 1

  export_filename = f'employee_matrix_{start_date_str}_to_{end_date_str}.xlsx'
  export_path = os.path.join(app.config['UPLOAD_FOLDER'], export_filename)
  wb.save(export_path)
  return send_file(export_path, as_attachment=True)

@app.route('/shutdown')
def shutdown():
  pwd = request.args.get('pwd', '')
  admin_pass = os.getenv('ADMIN_PWD', 'Gamek@789')
  dev_pass = os.getenv('DEV_PWD', 'Shama@8577')
  if pwd == admin_pass or pwd == dev_pass:
    func = request.environ.get('werkzeug.server.shutdown')
    if func:
      func()
    return 'Server successfully shutting down...'
  return 'Invalid password for shutdown!', 403

if __name__ == '__main__':
  app.run(host='0.0.0.0', port=5000, debug=True)
